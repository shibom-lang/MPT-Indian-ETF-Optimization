import yfinance as yf
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill

TICKERS = {
    'NIFTYBEES.NS': 'Nifty 50 BeES',
    'JUNIORBEES.NS': 'Junior BeES',
    'BANKBEES.NS': 'Bank BeES',
    'MID150BEES.NS': 'Midcap 150 BeES',
    'MON100.NS': 'Motilal Nasdaq 100',
    'GOLDBEES.NS': 'Gold BeES',
    'SILVERBEES.NS': 'Silver BeES',
    'SETF10GILT.NS': '10-Yr G-Sec ETF',
    'LIQUIDBEES.NS': 'Liquid BeES'
}

STRATEGY_STATS = {
    'maxsharpe': {'return': 0.1944, 'volatility': 0.1302, 'weights': {
        'GOLDBEES.NS': 0.25, 'MID150BEES.NS': 0.25, 'MON100.NS': 0.25, 'JUNIORBEES.NS': 0.15, 'NIFTYBEES.NS': 0.10
    }}
}

def fetch_data():
    raw = yf.download(list(TICKERS.keys()), period='5y', auto_adjust=True, progress=False, group_by='column')
    closes = raw['Close'] if isinstance(raw.columns, pd.MultiIndex) else raw
    closes.dropna(how='all', inplace=True)
    closes.ffill(inplace=True)
    
    benchmark = closes['NIFTYBEES.NS']
    benchmark_ret = benchmark.pct_change().dropna()
    
    stats = {}
    for t in TICKERS.keys():
        prices = closes[t].dropna()
        if len(prices) == 0:
            prices = pd.Series([100]*len(benchmark), index=benchmark.index)
        
        current_price = float(prices.iloc[-1])
        ret = prices.pct_change().dropna()
        
        # Align series to avoid errors
        align_ret, align_bench = ret.align(benchmark_ret, join='inner')
        cov = align_ret.cov(align_bench)
        var = align_bench.var()
        beta = cov / var if var != 0 else 1.0
        
        div_yield = 0.015 if 'LIQUID' not in t and 'GSEC' not in t else 0.06
        if 'GOLD' in t or 'SILVER' in t:
            div_yield = 0.0
            
        stats[t] = {
            'price': current_price,
            'beta': beta,
            'div_yield': div_yield
        }
    return stats

def create_excel(amount, strategy_id):
    stats = fetch_data()
    strategy = STRATEGY_STATS['maxsharpe']
    weights = strategy['weights']
    
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Portfolio Beta Calculator"
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    headers = ["Ticker Symbol", "Company Name", "Shares Held", "Share Price", "Portfolio Value", "Weight", "5Y Beta", "Weighted Beta"]
    ws1.append(headers)
    for col in range(1, len(headers)+1):
        cell = ws1.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        
    for t, w in weights.items():
        if w > 0:
            val = amount * w
            price = stats[t]['price']
            shares = int(val / price)
            actual_val = shares * price
            beta = stats[t]['beta']
            w_beta = beta * w
            ws1.append([t, TICKERS[t], shares, price, actual_val, w, beta, w_beta])
            
    wb.save("/tmp/test.xlsx")
    print("Excel generated at /tmp/test.xlsx")

create_excel(100000, 'maxsharpe')
